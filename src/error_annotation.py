"""Human error-cause annotation of the 512-token grid error samples.

The samples (results/tables/grid_max_length_512/error_analysis_sample_{task}.csv) hold misclassified
test items of BanglaBERT full fine-tuning (seed 42), stratified by confidence within each condition.
Two annotators code the primary cause of each error independently; the shift condition is hidden and
the item order is shuffled. Disagreements can optionally be resolved by discussion in an adjudication
workbook; without it, the cause distribution is reported per annotator and for the agreed items.

    python -m src.error_annotation prepare        one workbook per annotator
    python -m src.error_annotation adjudication   workbook listing the disagreements (optional)
    python -m src.error_annotation score          agreement and cause distributions
"""

import argparse
from pathlib import Path

import numpy as np
import pandas as pd
from sklearn.metrics import cohen_kappa_score

from src.config import RESULTS_DIR

TASKS = ("offensive", "sentiment")
SAMPLE_PATH = RESULTS_DIR / "tables" / "grid_max_length_512" / "error_analysis_sample_{}.csv"
AUDIT_DIR = Path(__file__).resolve().parent.parent / "data" / "annotation"
KEY_PATH = AUDIT_DIR / "error_annotation_key.csv"
WORKBOOK = "error_annotation_annotator_{}.xlsx"
ADJUDICATION = AUDIT_DIR / "error_annotation_adjudication.xlsx"
OUT_DIR = RESULTS_DIR / "tables" / "grid_max_length_512"
ORDER_SEED = 2027
B_BOOTSTRAP = 10_000

CAUSES = [
    "gold label disputed",
    "dialect or non-standard lexicon",
    "code-mixing or script",
    "surface corruption",
    "implicit meaning",
    "negation or contrast",
    "mixed polarity or target",
    "no identifiable cause",
]

INSTRUCTIONS = [
    "Error-cause annotation: instructions",
    "",
    "Every row is a text that the model misclassified. 'gold_label' is the dataset label, 'predicted_label' the",
    "model's answer. Choose the ONE primary reason why a model would get this text wrong (drop-down in 'cause').",
    "Work alone, without machine translation or AI tools, and do not discuss items before both workbooks are done.",
    "",
    "Apply the categories in this order and take the first that fits:",
    "1. gold label disputed: you would not give the gold label yourself (wrong or very doubtful annotation).",
    "2. dialect or non-standard lexicon: the words that decide the label are dialectal, regional, slang, or in an",
    "   unusual spelling of Bangla.",
    "3. code-mixing or script: the words that decide the label are English, Hindi, or romanised Bangla/Hindi.",
    "4. surface corruption: typing errors or character noise damage the words that decide the label.",
    "5. implicit meaning: sarcasm, irony, rhetorical question, insult without explicit abusive words, or world",
    "   knowledge needed.",
    "6. negation or contrast: a negation or a 'but'/concession clause reverses the meaning.",
    "7. mixed polarity or target: the text contains both polarities or several targets and the model followed",
    "   the wrong part.",
    "8. no identifiable cause: the text is clear, standard and explicit, yet the model failed.",
    "",
    "Use 'note' only when it helps (for example, the decisive word, or why the gold label is doubtful).",
]


def short_id(task: str, pos: int) -> str:
    return f"{task[:3]}_e{pos:03d}"


def prepare() -> pd.DataFrame:
    rng = np.random.default_rng(ORDER_SEED)
    rows = []
    for task in TASKS:
        df = pd.read_csv(str(SAMPLE_PATH).format(task), encoding="utf-8-sig")
        for pos, j in enumerate(rng.permutation(len(df)), start=1):
            r = df.iloc[j]
            rows.append({"item_id": short_id(task, pos), "task": task, "error_id": r["error_id"],
                         "condition": r["condition"], "confidence": r["confidence"], "text": str(r["text"]),
                         "gold_label": r["true_label"], "predicted_label": r["predicted_label"]})
    return pd.DataFrame(rows)


def write_workbook(key: pd.DataFrame, path: Path, extra_columns: list = None, rows: pd.DataFrame = None) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.worksheet.datavalidation import DataValidation

    rows = key if rows is None else rows
    extra_columns = extra_columns or []
    header = ["item_id", "text", "gold_label", "predicted_label", *extra_columns, "cause", "note"]
    wb = Workbook()
    ws = wb.active
    ws.title = "instructions"
    for r, line in enumerate(INSTRUCTIONS, start=1):
        ws.cell(row=r, column=1, value=line).font = Font(bold=(r == 1), size=12 if r == 1 else 11)
    ws.column_dimensions["A"].width = 115
    cause_col = chr(ord("A") + header.index("cause"))
    for task in TASKS:
        items = rows[rows.task == task]
        sheet = wb.create_sheet(task)
        sheet.append(header)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for row in items.itertuples():
            sheet.append([row.item_id, row.text, row.gold_label, row.predicted_label,
                          *[getattr(row, c) for c in extra_columns], getattr(row, "cause", None), None])
        dv = DataValidation(type="list", formula1='"' + ",".join(CAUSES) + '"', allow_blank=True,
                            showErrorMessage=True, errorTitle="Invalid cause", error="Choose a value from the list.")
        sheet.add_data_validation(dv)
        dv.add(f"{cause_col}2:{cause_col}{len(items) + 1}")
        widths = {"A": 10, "B": 90, "C": 15, "D": 15}
        for i, c in enumerate(extra_columns):
            widths[chr(ord("E") + i)] = 28
        widths[cause_col] = 30
        widths[chr(ord(cause_col) + 1)] = 30
        for col, w in widths.items():
            sheet.column_dimensions[col].width = w
        for r in range(2, len(items) + 2):
            sheet.cell(row=r, column=2).alignment = Alignment(wrap_text=True, vertical="top")
        sheet.freeze_panes = "A2"
    wb.save(path)


def read_causes(path: Path, column: str = "cause") -> pd.DataFrame:
    frames = [pd.read_excel(path, sheet_name=task, dtype=str)[["item_id", column]] for task in TASKS]
    df = pd.concat(frames, ignore_index=True)
    df[column] = df[column].fillna("").str.strip()
    return df


def annotations() -> tuple:
    key = pd.read_csv(KEY_PATH)
    paths = sorted(AUDIT_DIR.glob(WORKBOOK.format("*")))
    if len(paths) != 2:
        raise SystemExit(f"Expected two annotator workbooks in {AUDIT_DIR}, found {len(paths)}")
    names = [p.stem.replace("error_annotation_", "") for p in paths]
    for path, name in zip(paths, names):
        key = key.merge(read_causes(path).rename(columns={"cause": name}), on="item_id", how="left")
    invalid = {n: int((~key[n].isin(CAUSES)).sum()) for n in names}
    if any(invalid.values()):
        print(f"Items without a valid cause: {invalid}")
    return key, names


def adjudication() -> None:
    if ADJUDICATION.exists():
        raise SystemExit(f"{ADJUDICATION.name} already exists and may hold resolved causes; rename it to rebuild")
    key, names = annotations()
    coded = key[names[0]].isin(CAUSES) & key[names[1]].isin(CAUSES)
    disagree = key[(key[names[0]] != key[names[1]]) | ~coded].copy()
    disagree["cause"] = None
    write_workbook(key, ADJUDICATION, extra_columns=names, rows=disagree)
    print(f"Wrote {ADJUDICATION}: {len(disagree)} items to resolve ({int((~coded).sum())} not coded by both)")


def bootstrap_kappa(a: np.ndarray, b: np.ndarray, rng) -> tuple:
    stats = []
    for _ in range(B_BOOTSTRAP):
        i = rng.integers(0, len(a), len(a))
        stats.append(cohen_kappa_score(a[i], b[i]))
    return float(np.nanquantile(stats, 0.025)), float(np.nanquantile(stats, 0.975))


def gwet_ac1(a: np.ndarray, b: np.ndarray) -> float:
    pi = np.array([((a == c).mean() + (b == c).mean()) / 2 for c in CAUSES])
    pe = float((pi * (1 - pi)).sum() / (len(CAUSES) - 1))
    return float(((a == b).mean() - pe) / (1 - pe))


def score() -> None:
    key, names = annotations()
    if ADJUDICATION.exists():
        final = read_causes(ADJUDICATION).rename(columns={"cause": "adjudicated"})
        key = key.merge(final, on="item_id", how="left")
    else:
        key["adjudicated"] = ""
    coded = key[names[0]].isin(CAUSES) & key[names[1]].isin(CAUSES)
    agreed = coded & (key[names[0]] == key[names[1]])
    key["final_cause"] = np.where(agreed, key[names[0]], key["adjudicated"].fillna(""))
    rng = np.random.default_rng(0)
    rows = []
    for task, full in key.groupby("task"):
        g = full[full[names[0]].isin(CAUSES) & full[names[1]].isin(CAUSES)]
        a, b = g[names[0]].to_numpy(), g[names[1]].to_numpy()
        lo, hi = bootstrap_kappa(a, b, rng)
        rows.append({"task": task, "condition": "all", "n_items": len(g), "n_not_double_coded": len(full) - len(g),
                     "agreement": float((a == b).mean()), "cohen_kappa": float(cohen_kappa_score(a, b)),
                     "kappa_ci_low": lo, "kappa_ci_high": hi, "gwet_ac1": gwet_ac1(a, b),
                     "n_disagreed_or_uncoded": int((~agreed[full.index]).sum()),
                     "n_unresolved": int((~full["final_cause"].isin(CAUSES)).sum())})
        for cond, h in g.groupby("condition"):
            ha, hb = h[names[0]].to_numpy(), h[names[1]].to_numpy()
            rows.append({"task": task, "condition": cond, "n_items": len(h),
                         "n_not_double_coded": int((full.condition == cond).sum()) - len(h),
                         "agreement": float((ha == hb).mean()), "cohen_kappa": float(cohen_kappa_score(ha, hb)),
                         "gwet_ac1": gwet_ac1(ha, hb)})
    agreement = pd.DataFrame(rows)
    for c in ("n_disagreed_or_uncoded", "n_unresolved"):
        agreement[c] = agreement[c].astype("Int64")
    # One row per (task, condition, coding, cause). Shares of each annotator are over the items that annotator
    # coded; "agreed" and "adjudicated" shares are over all items of the condition.
    key["agreed"] = np.where(agreed, key[names[0]], "")
    codings = [*names, "agreed"] + (["adjudicated"] if ADJUDICATION.exists() else [])
    if "adjudicated" in codings:
        key["adjudicated"] = key["final_cause"]
    frames = []
    for coding in codings:
        counts = (key.groupby(["task", "condition"])[coding].value_counts().rename("n_items").reset_index()
                  .rename(columns={coding: "cause"}))
        totals = key.groupby(["task", "condition"])[coding].agg(
            lambda s: s.isin(CAUSES).sum() if coding in names else len(s)).rename("denominator").reset_index()
        counts = counts[counts.cause.isin(CAUSES)].merge(totals, on=["task", "condition"])
        counts.insert(2, "coding", coding)
        frames.append(counts)
    distribution = pd.concat(frames, ignore_index=True)
    distribution["share"] = distribution.n_items / distribution.denominator
    distribution = distribution.drop(columns="denominator").sort_values(["task", "condition", "coding", "cause"])
    agreement.to_csv(OUT_DIR / "error_annotation_agreement.csv", index=False, float_format="%.4f", lineterminator="\n")
    distribution.to_csv(OUT_DIR / "error_annotation_cause_distribution.csv", index=False, float_format="%.4f",
                        lineterminator="\n")
    pd.set_option("display.width", 200)
    print(agreement.round(3).to_string(index=False))
    print(distribution.pivot_table(index=["task", "condition", "coding"], columns="cause", values="n_items",
                                   fill_value=0).astype(int).to_string())


def main():
    parser = argparse.ArgumentParser(description="Human error-cause annotation of the 512-token grid error samples.")
    parser.add_argument("action", choices=["prepare", "adjudication", "score"])
    args = parser.parse_args()
    if args.action == "prepare":
        AUDIT_DIR.mkdir(parents=True, exist_ok=True)
        key = prepare()
        key.to_csv(KEY_PATH, index=False, lineterminator="\n")
        for i in (1, 2):
            path = AUDIT_DIR / WORKBOOK.format(i)
            if path.exists():
                print(f"Kept existing {path.name} (never overwritten)")
            else:
                write_workbook(key, path)
        print(f"Wrote {KEY_PATH} ({len(key)} items); workbooks in {AUDIT_DIR}")
        print(key.groupby(["task", "condition"]).size().to_string())
    elif args.action == "adjudication":
        adjudication()
    else:
        score()


if __name__ == "__main__":
    main()
