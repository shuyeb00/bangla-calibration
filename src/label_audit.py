"""Blind human relabelling of in-domain and code-mixed test items (label-definition audit).

Offensive D2 models lose all discrimination (ranking AUROC about 0.5). This audit separates two
explanations: the models fail on code-mixed input (covariate shift), or the OffMix-3L / SentMix-3L
labels follow a different label definition than BD-SHS / SentNoB (concept shift). Two annotators
label random D0 and D2 test items without seeing the gold label, the split or the model output;
agreement with gold on D0 is the reference for agreement with gold on D2.

    python -m src.label_audit sample      writes the answer key and one workbook per annotator
    python -m src.label_audit score       reads the filled workbooks, writes the agreement table
"""

import argparse
from pathlib import Path

import numpy as np
import pandas as pd
from sklearn.metrics import cohen_kappa_score

from src.config import DATA_DIR, RESULTS_DIR

AUDIT_DIR = Path(__file__).resolve().parent.parent / "data" / "annotation"
KEY_PATH = AUDIT_DIR / "label_audit_key.csv"
WORKBOOK = "label_audit_annotator_{}.xlsx"
OUT_PATH = RESULTS_DIR / "tables" / "label_audit_agreement.csv"
SAMPLE_SEED = 2026
N_ITEMS = {"D0": 60, "D2": 100}
SOURCES = {"D0": "test.csv", "D2": "d2_codemixed_test.csv"}
LABELS = {
    "offensive": {0: "Not offensive", 1: "Offensive"},
    "sentiment": {0: "Negative", 1: "Neutral", 2: "Positive"},
}
CANNOT_JUDGE = "Cannot judge"
B_BOOTSTRAP = 10_000

INSTRUCTIONS = [
    "Label audit: instructions",
    "",
    "Work alone. Do not discuss items with the other annotator until both workbooks are submitted.",
    "Do not use machine translation, search engines or AI tools. Read each text once or twice and decide.",
    "Fill only the 'label' column (drop-down) and, if useful, the 'note' column. Do not change other cells.",
    "",
    "Sheet 'offensive': choose 'Offensive' if the text insults, demeans, threatens or uses profane or hateful",
    "language towards a person or group (including the person addressed). Criticism without insult is",
    "'Not offensive'.",
    "Sheet 'sentiment': choose the overall attitude of the writer. 'Positive' or 'Negative' if one attitude",
    "dominates; 'Neutral' if the text is factual, has no clear attitude, or is balanced.",
    "",
    "Choose 'Cannot judge' only if you cannot understand the part of the text that decides the label (for",
    "example, a script or language you do not read). Do not use it for hard but readable cases: give your",
    "best judgement instead.",
    "Expected time: about 15 seconds per item, 320 items in total. Save the file as it is (.xlsx).",
]


def sample_items() -> pd.DataFrame:
    rng = np.random.default_rng(SAMPLE_SEED)
    rows = []
    for task in LABELS:
        task_rows = []
        for split, n in N_ITEMS.items():
            df = pd.read_csv(DATA_DIR / task / SOURCES[split])
            idx = np.sort(rng.choice(len(df), size=n, replace=False))
            for i in idx:
                task_rows.append({"task": task, "split": split, "source_file": SOURCES[split], "source_row": int(i),
                                  "text": str(df["text_clean"].iloc[i]), "gold": LABELS[task][int(df["labels"].iloc[i])]})
        order = rng.permutation(len(task_rows))
        for pos, j in enumerate(order, start=1):
            rows.append({**task_rows[j], "item_id": f"{task[:3]}_{pos:03d}"})
    cols = ["item_id", "task", "split", "source_file", "source_row", "gold", "text"]
    return pd.DataFrame(rows)[cols]


def write_workbook(key: pd.DataFrame, path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "instructions"
    for r, line in enumerate(INSTRUCTIONS, start=1):
        ws.cell(row=r, column=1, value=line).font = Font(bold=(r == 1), size=12 if r == 1 else 11)
    ws.column_dimensions["A"].width = 110
    for task, names in LABELS.items():
        sheet = wb.create_sheet(task)
        sheet.append(["item_id", "text", "label", "note"])
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        items = key[key.task == task]
        for row in items.itertuples():
            sheet.append([row.item_id, row.text, None, None])
        options = ",".join(list(names.values()) + [CANNOT_JUDGE])
        dv = DataValidation(type="list", formula1=f'"{options}"', allow_blank=True, showErrorMessage=True,
                            errorTitle="Invalid label", error="Choose a value from the list.")
        sheet.add_data_validation(dv)
        dv.add(f"C2:C{len(items) + 1}")
        sheet.column_dimensions["A"].width = 10
        sheet.column_dimensions["B"].width = 100
        sheet.column_dimensions["C"].width = 16
        sheet.column_dimensions["D"].width = 30
        for r in range(2, len(items) + 2):
            sheet.cell(row=r, column=2).alignment = Alignment(wrap_text=True, vertical="top")
        sheet.freeze_panes = "A2"
    wb.save(path)


def read_workbook(path: Path) -> pd.DataFrame:
    frames = [pd.read_excel(path, sheet_name=task, dtype=str)[["item_id", "label"]] for task in LABELS]
    df = pd.concat(frames, ignore_index=True)
    df["label"] = df["label"].fillna("").str.strip()
    return df


def bootstrap_difference(a: np.ndarray, b: np.ndarray, rng) -> tuple:
    diffs = np.array([rng.choice(a, len(a)).mean() - rng.choice(b, len(b)).mean() for _ in range(B_BOOTSTRAP)])
    return float(np.quantile(diffs, 0.025)), float(np.quantile(diffs, 0.975))


def score() -> pd.DataFrame:
    key = pd.read_csv(KEY_PATH, dtype={"source_row": int})
    annotators = sorted(AUDIT_DIR.glob(WORKBOOK.format("*")))
    if len(annotators) < 2:
        raise SystemExit(f"Expected two filled workbooks in {AUDIT_DIR}")
    merged = key.copy()
    names = []
    for path in annotators:
        name = path.stem.replace("label_audit_", "")
        labels = read_workbook(path).rename(columns={"label": name})
        merged = merged.merge(labels, on="item_id", how="left")
        names.append(name)
    missing = merged[names].isin(["", np.nan]).sum()
    if missing.any():
        print(f"Unlabelled items per annotator: {missing.to_dict()}")

    rng = np.random.default_rng(0)
    rows = []
    for (task, split), g in merged.groupby(["task", "split"], sort=False):
        row = {"task": task, "split": split, "n_items": len(g)}
        judged_all = np.ones(len(g), dtype=bool)
        for name in names:
            judged = g[name].isin(LABELS[task].values()).to_numpy()
            judged_all &= judged
            row[f"{name}_cannot_judge_rate"] = float((g[name] == CANNOT_JUDGE).mean())
            row[f"{name}_agreement_with_gold"] = float((g[name][judged] == g["gold"][judged]).mean())
            row[f"{name}_kappa_with_gold"] = float(cohen_kappa_score(g["gold"][judged], g[name][judged]))
        both = g[judged_all]
        row["n_judged_by_both"] = len(both)
        row["inter_annotator_agreement"] = float((both[names[0]] == both[names[1]]).mean())
        row["inter_annotator_kappa"] = float(cohen_kappa_score(both[names[0]], both[names[1]]))
        consensus = both[both[names[0]] == both[names[1]]]
        row["n_consensus"] = len(consensus)
        row["consensus_agreement_with_gold"] = float((consensus[names[0]] == consensus["gold"]).mean())
        rows.append(row)
    table = pd.DataFrame(rows)
    # Item-level score: share of the annotators who judged the item that agree with the gold label.
    judged = pd.concat([merged[n].isin(set(LABELS["offensive"].values()) | set(LABELS["sentiment"].values()))
                        for n in names], axis=1).to_numpy()
    agree = pd.concat([merged[n] == merged["gold"] for n in names], axis=1).to_numpy() & judged
    merged["gold_agreement"] = np.where(judged.any(axis=1), agree.sum(axis=1) / np.maximum(judged.sum(axis=1), 1),
                                        np.nan)
    for task in LABELS:
        sub = merged[(merged.task == task) & merged.gold_agreement.notna()]
        d0 = sub[sub.split == "D0"].gold_agreement.to_numpy()
        d2 = sub[sub.split == "D2"].gold_agreement.to_numpy()
        if len(d0) and len(d2):
            lo, hi = bootstrap_difference(d0, d2, rng)
            mask = (table.task == task) & (table.split == "D2")
            table.loc[mask, "gold_agreement_drop_vs_d0"] = d0.mean() - d2.mean()
            table.loc[mask, "gold_agreement_drop_ci_low"] = lo
            table.loc[mask, "gold_agreement_drop_ci_high"] = hi
    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    table.to_csv(OUT_PATH, index=False, float_format="%.4f", lineterminator="\n")
    return table


def main():
    parser = argparse.ArgumentParser(description="Blind relabelling audit of D0 and D2 test items.")
    parser.add_argument("action", choices=["sample", "score"])
    parser.add_argument("--annotators", type=int, default=2)
    args = parser.parse_args()
    if args.action == "sample":
        AUDIT_DIR.mkdir(parents=True, exist_ok=True)
        key = sample_items()
        key.to_csv(KEY_PATH, index=False, lineterminator="\n")
        for i in range(1, args.annotators + 1):
            write_workbook(key, AUDIT_DIR / WORKBOOK.format(i))
        print(f"Wrote {KEY_PATH} ({len(key)} items) and {args.annotators} workbooks in {AUDIT_DIR}")
        print(key.groupby(["task", "split"]).gold.value_counts().unstack().to_string())
    else:
        pd.set_option("display.width", 220)
        print(score().round(3).T.to_string())
        print(f"Wrote {OUT_PATH}")


if __name__ == "__main__":
    main()
